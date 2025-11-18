#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""创建PPT数据Excel文件"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def create_excel_data():
    """创建包含第23页数据的Excel文件"""

    wb = Workbook()

    # 删除默认sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])

    # 创建第23页数据
    ws = wb.create_sheet("第23页_美股上涨驱动")

    # 设置标题行样式
    header_fill = PatternFill(start_color="002960", end_color="002960", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=12)

    # 时间轴数据
    ws['A1'] = '时间轴节点'
    ws['B1'] = '年份'
    ws['C1'] = '阶段名称'
    ws['D1'] = '描述'

    timeline_data = [
        ['节点1', '1933', '证券法体系', '制度红利'],
        ['节点2', '1971', '布雷顿森林体系瓦解', '美元霸权'],
        ['节点3', '1995', '互联网泡沫 + PC时代', '科技创新'],
        ['节点4', '2009', '金融危机后QE时代', '货币扩张'],
        ['节点5', '2020', '疫情后无限QE', '科技+货币'],
    ]

    for i, row in enumerate(timeline_data, start=2):
        ws[f'A{i}'] = row[0]
        ws[f'B{i}'] = row[1]
        ws[f'C{i}'] = row[2]
        ws[f'D{i}'] = row[3]

    # 制度红利数据
    ws['F1'] = '制度红利(1933-1940)'
    ws['F2'] = '1933证券法: 信息披露制度'
    ws['F3'] = '1934交易所法: SEC成立'
    ws['F4'] = '建立投资者信心'

    # 科技创新周期数据
    ws['H1'] = '科技创新周期'
    ws['H2'] = '1995-2000: 互联网革命(思科/微软)'
    ws['H3'] = '2007-2020: 移动互联网(苹果/FAANG)'
    ws['H4'] = '2020-现在: AI革命(Nvidia/微软/谷歌)'
    ws['H5'] = 'FAANG+微软市值占标普500超30%'

    # 货币政策数据
    ws['J1'] = '货币政策'
    ws['J2'] = '2009-2014: QE1-QE3推动估值扩张'
    ws['J3'] = '2020: 无限QE+零利率'
    ws['J4'] = '标普500 PE从15倍→25倍(估值扩张贡献40%涨幅)'

    # 洞察
    ws['L1'] = '核心洞察'
    ws['L2'] = '美股牛市=制度信任×科技创新×流动性支持'

    # 应用标题样式
    for col in ['A', 'F', 'H', 'J', 'L']:
        ws[f'{col}1'].fill = header_fill
        ws[f'{col}1'].font = header_font
        ws[f'{col}1'].alignment = Alignment(horizontal='center', vertical='center')

    # 调整列宽
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['H'].width = 35
    ws.column_dimensions['J'].width = 35
    ws.column_dimensions['L'].width = 40

    # 保存文件
    wb.save('ppt_data.xlsx')
    print("✓ Excel数据文件创建成功: ppt_data.xlsx")

if __name__ == '__main__':
    create_excel_data()
