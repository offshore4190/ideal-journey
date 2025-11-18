#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""美股牛市分析PPT生成器 - 第23页"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import openpyxl

# ============= 配色方案 =============
PRIMARY_BLUE = RGBColor(0, 41, 96)      # #002960 深蓝
SECONDARY_BLUE = RGBColor(0, 101, 189)  # #0065BD 中蓝
LIGHT_BLUE = RGBColor(201, 240, 255)    # #C9F0FF 浅蓝
IKEA_YELLOW = RGBColor(255, 219, 0)     # #FFDB00 黄色
WHITE = RGBColor(255, 255, 255)         # #FFFFFF 白色
GRAY = RGBColor(128, 128, 128)          # #808080 灰色

# ============= 排版规范 =============
SLIDE_WIDTH = Inches(10.0)              # 16:9
SLIDE_HEIGHT = Inches(5.625)
CONTENT_LEFT = Inches(0.8)              # 左边距
CONTENT_WIDTH = Inches(8.4)             # 内容区宽度
CONTENT_TOP = Inches(1.0)               # 上边距(标题下)

# 字体规范
TITLE_FONT_SIZE = Pt(18)                # 页面标题
BODY_FONT_SIZE = Pt(11)                 # 正文
CAPTION_FONT_SIZE = Pt(9)               # 注释/来源
MIN_FONT_SIZE = Pt(7)                   # 最小字号


class PPTGenerator:
    """PPT生成器"""

    def __init__(self, excel_file='ppt_data.xlsx'):
        """初始化"""
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_WIDTH
        self.prs.slide_height = SLIDE_HEIGHT
        self.excel_file = excel_file

    def read_excel_data(self, sheet_name):
        """读取Excel数据"""
        wb = openpyxl.load_workbook(self.excel_file)
        ws = wb[sheet_name]
        data = {}

        # 读取时间轴数据
        data['timeline'] = []
        for row in range(2, 7):  # 5个时间节点
            data['timeline'].append({
                'year': ws[f'B{row}'].value,
                'stage': ws[f'C{row}'].value,
                'desc': ws[f'D{row}'].value
            })

        # 读取内容框数据
        data['institution'] = [
            ws['F2'].value,
            ws['F3'].value,
            ws['F4'].value
        ]

        data['technology'] = [
            ws['H2'].value,
            ws['H3'].value,
            ws['H4'].value,
            ws['H5'].value
        ]

        data['monetary'] = [
            ws['J2'].value,
            ws['J3'].value,
            ws['J4'].value
        ]

        data['insight'] = ws['L2'].value

        wb.close()
        return data

    def add_title(self, slide, title_text):
        """添加标题"""
        title_box = slide.shapes.add_textbox(
            CONTENT_LEFT, Inches(0.4),
            CONTENT_WIDTH, Inches(0.5)
        )
        text_frame = title_box.text_frame
        text_frame.text = title_text

        # 标题样式
        para = text_frame.paragraphs[0]
        para.font.size = TITLE_FONT_SIZE
        para.font.bold = True
        para.font.color.rgb = PRIMARY_BLUE
        para.alignment = PP_ALIGN.LEFT

    def add_timeline(self, slide, timeline_data):
        """添加时间轴"""
        # 时间轴线条
        line_top = Inches(1.5)
        line_left = Inches(1.2)
        line_width = Inches(7.6)

        line = slide.shapes.add_connector(
            1,  # MSO_CONNECTOR_TYPE.STRAIGHT
            line_left, line_top,
            line_left + line_width, line_top
        )
        line.line.color.rgb = SECONDARY_BLUE
        line.line.width = Pt(3)

        # 添加时间节点
        node_spacing = line_width / 4  # 5个节点，4个间隔
        for i, node in enumerate(timeline_data):
            x_pos = line_left + (i * node_spacing)

            # 节点圆圈
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                x_pos - Inches(0.15), line_top - Inches(0.15),
                Inches(0.3), Inches(0.3)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = SECONDARY_BLUE
            circle.line.color.rgb = SECONDARY_BLUE

            # 年份（节点上方）
            year_box = slide.shapes.add_textbox(
                x_pos - Inches(0.3), line_top - Inches(0.5),
                Inches(0.6), Inches(0.25)
            )
            year_frame = year_box.text_frame
            year_frame.text = str(node['year'])
            year_para = year_frame.paragraphs[0]
            year_para.font.size = Pt(10)
            year_para.font.bold = True
            year_para.font.color.rgb = SECONDARY_BLUE
            year_para.alignment = PP_ALIGN.CENTER

            # 阶段名称（节点下方第1行）
            stage_box = slide.shapes.add_textbox(
                x_pos - Inches(0.5), line_top + Inches(0.25),
                Inches(1.0), Inches(0.3)
            )
            stage_frame = stage_box.text_frame
            stage_frame.text = node['stage']
            stage_frame.word_wrap = True
            stage_para = stage_frame.paragraphs[0]
            stage_para.font.size = Pt(8)
            stage_para.font.color.rgb = PRIMARY_BLUE
            stage_para.alignment = PP_ALIGN.CENTER

            # 描述（节点下方第2行）
            desc_box = slide.shapes.add_textbox(
                x_pos - Inches(0.5), line_top + Inches(0.55),
                Inches(1.0), Inches(0.25)
            )
            desc_frame = desc_box.text_frame
            desc_frame.text = node['desc']
            desc_para = desc_frame.paragraphs[0]
            desc_para.font.size = Pt(8)
            desc_para.font.bold = True
            desc_para.font.color.rgb = GRAY
            desc_para.alignment = PP_ALIGN.CENTER

    def add_content_box(self, slide, left, top, width, height, title, content_list, number):
        """添加内容框"""
        # 背景框
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, height
        )
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_BLUE
        box.line.color.rgb = SECONDARY_BLUE
        box.line.width = Pt(1.5)

        # 序号圆圈
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            left + Inches(0.15), top + Inches(0.1),
            Inches(0.25), Inches(0.25)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = SECONDARY_BLUE
        circle.line.color.rgb = SECONDARY_BLUE

        # 序号文字
        number_box = slide.shapes.add_textbox(
            left + Inches(0.15), top + Inches(0.1),
            Inches(0.25), Inches(0.25)
        )
        number_frame = number_box.text_frame
        number_frame.text = str(number)
        number_para = number_frame.paragraphs[0]
        number_para.font.size = Pt(12)
        number_para.font.bold = True
        number_para.font.color.rgb = WHITE
        number_para.alignment = PP_ALIGN.CENTER
        number_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        # 标题
        title_box = slide.shapes.add_textbox(
            left + Inches(0.5), top + Inches(0.1),
            width - Inches(0.6), Inches(0.3)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(11)
        title_para.font.bold = True
        title_para.font.color.rgb = PRIMARY_BLUE

        # 内容列表
        content_box = slide.shapes.add_textbox(
            left + Inches(0.2), top + Inches(0.45),
            width - Inches(0.4), height - Inches(0.55)
        )
        content_frame = content_box.text_frame
        content_frame.word_wrap = True

        for item in content_list:
            p = content_frame.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(9)
            p.font.color.rgb = PRIMARY_BLUE
            p.space_before = Pt(3)
            p.space_after = Pt(3)

        # 删除第一个空段落
        if content_frame.paragraphs[0].text == '':
            content_frame.paragraphs[0]._element.getparent().remove(
                content_frame.paragraphs[0]._element
            )

    def add_insight_box(self, slide, insight_text):
        """添加洞察框"""
        left = Inches(2.5)
        top = Inches(4.8)
        width = Inches(5.0)
        height = Inches(0.5)

        # 背景框
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, width, height
        )
        box.fill.solid()
        box.fill.fore_color.rgb = IKEA_YELLOW
        box.line.color.rgb = SECONDARY_BLUE
        box.line.width = Pt(2)

        # 文字
        text_box = slide.shapes.add_textbox(
            left + Inches(0.2), top + Inches(0.05),
            width - Inches(0.4), height - Inches(0.1)
        )
        text_frame = text_box.text_frame
        text_frame.text = f"💡 {insight_text}"
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        para = text_frame.paragraphs[0]
        para.font.size = Pt(11)
        para.font.bold = True
        para.font.color.rgb = PRIMARY_BLUE
        para.alignment = PP_ALIGN.CENTER

    def create_slide_23(self):
        """创建第23页：美股11轮牛市——制度红利+科技周期+货币宽松"""
        # 读取数据
        data = self.read_excel_data('第23页_美股上涨驱动')

        # 创建空白幻灯片
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)

        # 设置白色背景
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = WHITE

        # 添加标题
        self.add_title(slide, "第23页: 美股11轮牛市——制度红利+科技周期+货币宽松")

        # 添加时间轴
        self.add_timeline(slide, data['timeline'])

        # 添加三个内容框
        box_width = Inches(2.6)
        box_height = Inches(1.8)
        box_top = Inches(2.5)

        # 制度红利框
        self.add_content_box(
            slide,
            Inches(1.0), box_top,
            box_width, box_height,
            "制度红利(1933-1940)",
            data['institution'],
            1
        )

        # 科技创新周期框
        self.add_content_box(
            slide,
            Inches(3.8), box_top,
            box_width, box_height,
            "科技创新周期",
            data['technology'],
            2
        )

        # 货币政策框
        self.add_content_box(
            slide,
            Inches(6.6), box_top,
            box_width, box_height,
            "货币政策",
            data['monetary'],
            3
        )

        # 添加洞察框
        self.add_insight_box(slide, data['insight'])

        print("✓ 第23页创建成功")

    def save(self, filename='美股牛市分析.pptx'):
        """保存PPT"""
        self.prs.save(filename)
        print(f"✓ PPT已保存: {filename}")


def main():
    """主函数"""
    print("开始生成PPT...")
    generator = PPTGenerator('ppt_data.xlsx')

    # 创建第23页
    generator.create_slide_23()

    # 保存
    generator.save('美股牛市分析_第23页.pptx')
    print("\n✅ PPT生成完成！")


if __name__ == '__main__':
    main()
